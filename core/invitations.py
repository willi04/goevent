"""
core/invitations.py — Logique métier des invitations digitales (événements privés).

Partagée par le back-office (api/invitations.py, offre accompagnée) et, plus
tard, par le tableau de bord organisateur (offre autonome) : aucune règle ici
ne dépend de qui appelle, seulement de l'événement et des droits vérifiés par
la route.

Règles clés
- Le pass d'un invité = un jeton aléatoire de 32 caractères (secrets), jamais
  dérivé d'un id. QR = URL du pass : {PUBLIC_SITE_URL}/i/<jeton>.
- Quota en PERSONNES = somme des packs actifs. Avant tout pack : import
  possible, envoi / scan / suivi bloqués. Activation refusée si le pack est
  plus petit que la liste déjà importée.
- Un pass peut couvrir 1 à 20 personnes : au scan, l'agent choisit combien
  entrent (verrou ligne : deux agents ne peuvent pas dépasser le total).
- Hors ligne : la liste téléchargée contient l'empreinte SHA-256 du jeton,
  jamais le jeton. Synchro idempotente (client_uuid), dépassements signalés
  en conflit (deux téléphones hors ligne ont accepté le même pass).
- Données personnelles : purge J+30 (nom anonymisé, téléphone et jeton effacés).
"""
import hashlib
import re
import secrets
import uuid as uuidlib
from datetime import datetime, timedelta, timezone
from typing import Optional
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from sqlalchemy import func
from sqlalchemy.orm import Session

from core.config import (
    INVITATION_HOUSE_PHONE, INVITATION_PACK_SIZES, INVITATION_PAYMENT_CHANNELS,
    INVITATION_RETENTION_DAYS, INVITATION_SERVICE_LEVELS, settings,
)
from core.security import hash_pin, log_security_event
from models.event import Event
from models.invitation import EventGuest, GuestCheckin, InvitationPack, PrivateEventSettings
from models.user import User

TZ = ZoneInfo("Africa/Bangui")
MAX_SEATS_PER_PASS = 20
MAX_IMPORT_ROWS = 2000
_PASS_RE = re.compile(r"/i/([A-Za-z0-9_-]{20,64})(?:[/?#]|$)")
_TOKEN_RE = re.compile(r"^[A-Za-z0-9_-]{20,64}$")

DEFAULT_MESSAGE = (
    "Bonjour {nom} 🌸\n"
    "{hote} a le plaisir de vous inviter à {evenement}, le {date} à {lieu}.\n"
    "Voici votre pass personnel ({places}) : {lien}\n"
    "Présentez-le simplement à l'entrée. Au plaisir de vous y voir !"
)
MESSAGE_VARS = ("nom", "hote", "evenement", "date", "lieu", "places", "table", "lien")

_JOURS = ("lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche")
_MOIS = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
         "septembre", "octobre", "novembre", "décembre")


# ── utilitaires ─────────────────────────────────────────────────────
def utcnow() -> datetime:
    return datetime.utcnow()


def naive_utc(dt: Optional[datetime]) -> Optional[datetime]:
    if dt is None:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


def local(dt: Optional[datetime]) -> Optional[datetime]:
    return dt.replace(tzinfo=timezone.utc).astimezone(TZ) if dt else None


def fmt_date_fr(dt: Optional[datetime]) -> str:
    """« samedi 14 novembre 2026 à 16h00 » (heure de Bangui)."""
    d = local(dt)
    if not d:
        return ""
    return "%s %d %s %d à %dh%02d" % (_JOURS[d.weekday()], d.day, _MOIS[d.month - 1], d.year, d.hour, d.minute)


def fmt_time(dt: Optional[datetime]) -> str:
    d = local(dt)
    return d.strftime("%Hh%M") if d else "?"


def new_token() -> str:
    return secrets.token_urlsafe(24)  # 32 caractères, 192 bits


def token_hash(token: str) -> str:
    return hashlib.sha256(token.encode()).hexdigest()


def pass_url(token: str) -> str:
    return "%s/i/%s" % (settings.PUBLIC_SITE_URL.rstrip("/"), token)


def track_url(token: str) -> str:
    return "%s/suivi/%s" % (settings.PUBLIC_SITE_URL.rstrip("/"), token)


def parse_pass_qr(text: str) -> Optional[str]:
    """Jeton d'invitation contenu dans un QR, ou None (= billet classique)."""
    if not text:
        return None
    m = _PASS_RE.search(text.strip())
    return m.group(1) if m else None


def normalize_phone(raw) -> Optional[str]:
    """Numéro → E.164. 8 chiffres = Centrafrique (+236). Lève ValueError si invalide."""
    if raw is None:
        return None
    s = re.sub(r"[\s.\-()/]", "", str(raw))
    if not s:
        return None
    if s.startswith("00"):
        s = "+" + s[2:]
    if s.startswith("+"):
        digits = s[1:]
    elif len(s) == 8:
        digits = "236" + s
    elif s.startswith("236") and len(s) == 11:
        digits = s
    else:
        raise ValueError("numéro invalide (8 chiffres ou format international +XXX…)")
    if not digits.isdigit() or not 8 <= len(digits) <= 15:
        raise ValueError("numéro invalide")
    if digits.startswith("236") and len(digits) != 11:
        raise ValueError("numéro centrafricain invalide (+236 suivi de 8 chiffres)")
    return "+" + digits


def channels(active_only: bool = False) -> list:
    active = {c.strip() for c in settings.INVITATION_CHANNELS_ACTIVE.split(",") if c.strip()}
    out = [dict(c, active=c["code"] in active) for c in INVITATION_PAYMENT_CHANNELS]
    return [c for c in out if c["active"]] if active_only else out


def places_label(n: int) -> str:
    return "%d place%s" % (n, "s" if n > 1 else "")


# ── compte propriétaire des événements accompagnés ──────────────────
def get_house_organizer(db: Session) -> User:
    """Compte organisateur interne « GoEvent Invitations », INACTIF : aucune connexion
    possible, il ne sert que de propriétaire technique (organizer_id)."""
    u = db.query(User).filter(User.phone_number == INVITATION_HOUSE_PHONE).first()
    if u:
        return u
    u = User(phone_number=INVITATION_HOUSE_PHONE, full_name="GoEvent Invitations", org_name="GoEvent",
             role="organizer", is_active=False, pin_hash=hash_pin(secrets.token_hex(16)))
    db.add(u)
    db.flush()
    return u


# ── lecture ─────────────────────────────────────────────────────────
def get_private_event(db: Session, event_id: int, lock: bool = False) -> Event:
    q = db.query(Event).filter(Event.id == event_id, Event.visibility == "private")
    if lock:
        q = q.with_for_update()
    ev = q.first()
    if not ev:
        raise HTTPException(404, "Événement privé introuvable")
    return ev


def get_settings(db: Session, event_id: int) -> PrivateEventSettings:
    s = db.query(PrivateEventSettings).filter(PrivateEventSettings.event_id == event_id).first()
    if not s:
        raise HTTPException(404, "Réglages de l'événement privé introuvables")
    return s


def quota(db: Session, event_id: int) -> int:
    return int(db.query(func.coalesce(func.sum(InvitationPack.size), 0))
               .filter(InvitationPack.event_id == event_id, InvitationPack.status == "actif").scalar() or 0)


def seats_listed(db: Session, event_id: int, exclude_guest: Optional[int] = None) -> int:
    q = db.query(func.coalesce(func.sum(EventGuest.seats), 0)).filter(
        EventGuest.event_id == event_id, EventGuest.revoked_at.is_(None))
    if exclude_guest:
        q = q.filter(EventGuest.id != exclude_guest)
    return int(q.scalar() or 0)


def active_service(db: Session, event_id: int) -> Optional[str]:
    """Niveau de service le plus élevé parmi les packs actifs."""
    svcs = {r[0] for r in db.query(InvitationPack.service).filter(
        InvitationPack.event_id == event_id, InvitationPack.status == "actif").all()}
    return "accompagne" if "accompagne" in svcs else ("autonome" if svcs else None)


def require_active(db: Session, event_id: int):
    if quota(db, event_id) <= 0:
        raise HTTPException(400, "Invitations non activées : aucun pack payé pour cet événement.")


def stats(db: Session, event_id: int) -> dict:
    row = db.query(
        func.coalesce(func.sum(EventGuest.seats), 0),
        func.coalesce(func.sum(EventGuest.seats_used), 0),
        func.count(EventGuest.id),
        func.count(EventGuest.sent_at),
        func.count(EventGuest.first_checkin_at),
    ).filter(EventGuest.event_id == event_id, EventGuest.revoked_at.is_(None)).one()
    expected, entered = int(row[0]), int(row[1])
    entered_shown = min(entered, expected) if expected else entered
    conflicts = db.query(func.count(GuestCheckin.id)).filter(
        GuestCheckin.event_id == event_id, GuestCheckin.conflict == True).scalar()  # noqa: E712
    return {
        "expected": expected, "entered": entered_shown, "remaining": max(0, expected - entered),
        "rate": round(entered_shown / expected * 100, 1) if expected else 0.0,
        "passes": int(row[2]), "passes_sent": int(row[3]), "passes_arrived": int(row[4]),
        "conflicts": int(conflicts or 0),
    }


# ── message WhatsApp ────────────────────────────────────────────────
def render_message(ev: Event, st: PrivateEventSettings, g: EventGuest) -> str:
    values = {
        "nom": g.full_name, "hote": st.host_label or st.client_name, "evenement": ev.title,
        "date": fmt_date_fr(ev.event_date), "lieu": ev.location or "", "places": places_label(g.seats),
        "table": g.table_label or "", "lien": pass_url(g.token) if g.token else "",
    }
    tpl = st.message_template or DEFAULT_MESSAGE
    return re.sub(r"\{(%s)\}" % "|".join(MESSAGE_VARS), lambda m: values[m.group(1)], tpl)


def wa_url(phone: Optional[str], text: str) -> Optional[str]:
    if not phone:
        return None
    return "https://wa.me/%s?text=%s" % (phone.lstrip("+"), quote(text, safe=""))


def check_template(tpl: Optional[str]) -> Optional[str]:
    if tpl is None:
        return None
    tpl = tpl.strip()
    if not tpl:
        return None
    if len(tpl) > 1000:
        raise HTTPException(400, "Message trop long (1000 caractères max).")
    if "{lien}" not in tpl:
        raise HTTPException(400, "Le message doit contenir {lien} (le pass personnel de l'invité).")
    unknown = set(re.findall(r"\{([a-z_]+)\}", tpl)) - set(MESSAGE_VARS)
    if unknown:
        raise HTTPException(400, "Variable inconnue : {%s}. Possibles : %s."
                            % (sorted(unknown)[0], ", ".join("{%s}" % v for v in MESSAGE_VARS)))
    return tpl


# ── import des invités ──────────────────────────────────────────────
def _clean(v, n):
    if v is None:
        return None
    v = "".join(ch for ch in str(v) if ord(ch) >= 32).strip()
    return v[:n] or None


def validate_rows(db: Session, event_id: int, rows: list) -> dict:
    """Valide les lignes (dicts) sans rien écrire. Retourne valid / errors / quota."""
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(400, "Import limité à %d lignes à la fois." % MAX_IMPORT_ROWS)
    existing = {p for (p,) in db.query(EventGuest.phone).filter(
        EventGuest.event_id == event_id, EventGuest.phone.isnot(None), EventGuest.revoked_at.is_(None)).all()}
    seen, valid, errors = set(), [], []
    for i, r in enumerate(rows, start=1):
        line = r.get("line") or i
        name = _clean(r.get("full_name"), 120)
        if not name or len(name) < 2:
            errors.append({"line": line, "message": "nom manquant"}); continue
        try:
            phone = normalize_phone(r.get("phone"))
        except ValueError as e:
            errors.append({"line": line, "name": name, "message": str(e)}); continue
        try:
            seats = int(str(r.get("seats") or 1).strip())
        except ValueError:
            errors.append({"line": line, "name": name, "message": "nombre de places invalide"}); continue
        if not 1 <= seats <= MAX_SEATS_PER_PASS:
            errors.append({"line": line, "name": name, "message": "places : entre 1 et %d" % MAX_SEATS_PER_PASS}); continue
        if phone and (phone in existing or phone in seen):
            errors.append({"line": line, "name": name, "message": "doublon (numéro déjà dans la liste)"}); continue
        if phone:
            seen.add(phone)
        valid.append({"line": line, "full_name": name, "phone": phone, "seats": seats,
                      "table_label": _clean(r.get("table_label"), 40), "category": _clean(r.get("category"), 40)})
    q, used = quota(db, event_id), seats_listed(db, event_id)
    adding = sum(v["seats"] for v in valid)
    return {"valid": valid, "errors": errors,
            "quota": {"pack": q, "listed": used, "adding": adding, "after": used + adding,
                      "exceeds": bool(q) and used + adding > q}}


def import_guests(db: Session, ev: Event, rows: list, user: User, dry_run: bool, request=None) -> dict:
    report = validate_rows(db, ev.id, rows)
    report["created"] = 0
    if dry_run:
        return report
    if report["quota"]["exceeds"]:
        raise HTTPException(400, "Le pack actif couvre %d personnes : cette liste en compterait %d. "
                                 "Activez un pack supérieur ou retirez des invités."
                            % (report["quota"]["pack"], report["quota"]["after"]))
    for v in report["valid"]:
        tok = new_token()
        db.add(EventGuest(event_id=ev.id, token=tok, token_hash=token_hash(tok), full_name=v["full_name"],
                          phone=v["phone"], seats=v["seats"], table_label=v["table_label"],
                          category=v["category"], created_by=user.id))
    db.commit()
    report["created"] = len(report["valid"])
    log_security_event(db, "invitation_guests_imported", request, user.id,
                       "event %d : %d pass, %d personnes, %d erreurs"
                       % (ev.id, report["created"], report["quota"]["adding"], len(report["errors"])))
    return report


# ── lecture de fichiers (.xlsx, .docx) ───────────────────────────────
# Lecture SIMPLE : le contenu est converti en lignes séparées par des
# tabulations, puis repasse par l'import habituel (mêmes colonnes, même
# vérification). Valeurs uniquement : formules recalculées ignorées, macros
# jamais exécutées. Le fichier est lu en mémoire puis jeté (rien n'est stocké).
MAX_FILE_BYTES = 2 * 1024 * 1024
_MAX_UNZIPPED = 40 * 1024 * 1024
_MAX_COLS = 10


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)                      # 75000001.0 → « 75000001 » (numéros saisis en nombres)
    if isinstance(v, datetime):
        v = v.strftime("%d/%m/%Y")
    return re.sub(r"[\t\r\n]+", " ", str(v)).strip()


def _check_zip(data: bytes):
    import io
    import zipfile
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        infos = z.infolist()
    except zipfile.BadZipFile:
        raise HTTPException(400, "Fichier illisible ou endommagé.")
    if len(infos) > 2000 or sum(i.file_size for i in infos) > _MAX_UNZIPPED:
        raise HTTPException(400, "Fichier trop volumineux une fois décompressé.")


def file_to_text(filename: str, data: bytes) -> dict:
    """Contenu d'un .xlsx (1re feuille) ou d'un .docx (tableaux, sinon lignes de texte) → texte tabulé."""
    import io
    name = (filename or "").lower()
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(400, "Fichier trop lourd (2 Mo maximum).")
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":   # ancien format binaire Office
        raise HTTPException(400, "Ancien format Word/Excel (.doc / .xls) : ouvrez le fichier et enregistrez-le "
                                 "en .docx ou .xlsx (Fichier → Enregistrer sous), ou copiez-collez la liste.")
    if data[:2] != b"PK":
        raise HTTPException(400, "Format non reconnu : fichiers acceptés .xlsx, .docx, .csv ou .txt.")
    _check_zip(data)
    lines = []
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(400, "Fichier Excel illisible.")
        try:
            ws = wb.worksheets[0]
            for row in ws.iter_rows(max_col=_MAX_COLS, values_only=True):
                cells = [_cell(v) for v in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip("\t"))
                if len(lines) > MAX_IMPORT_ROWS:
                    break
        finally:
            wb.close()
        fmt = "xlsx"
    elif name.endswith(".docx"):
        import docx
        try:
            d = docx.Document(io.BytesIO(data))
        except Exception:
            raise HTTPException(400, "Fichier Word illisible.")
        for table in d.tables:
            for row in table.rows:
                cells, seen = [], set()
                for c in row.cells[:_MAX_COLS]:
                    if id(c._tc) in seen:     # cellule fusionnée : une seule fois
                        continue
                    seen.add(id(c._tc))
                    cells.append(_cell(c.text))
                if any(cells):
                    lines.append("\t".join(cells).rstrip("\t"))
        if not lines:                        # pas de tableau : une ligne de texte = une ligne de la liste
            lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
        fmt = "docx"
    else:
        raise HTTPException(400, "Format non reconnu : fichiers acceptés .xlsx, .docx, .csv ou .txt.")
    if len(lines) > MAX_IMPORT_ROWS + 1:
        raise HTTPException(400, "Liste trop longue : %d lignes maximum par import." % MAX_IMPORT_ROWS)
    if not lines:
        raise HTTPException(400, "Aucune ligne trouvée dans le fichier.")
    return {"format": fmt, "rows": len(lines), "text": "\n".join(lines)}


# ── scan ─────────────────────────────────────────────────────────────
def guest_public(g: EventGuest) -> dict:
    return {"guest_name": g.full_name, "table": g.table_label, "category": g.category,
            "seats": g.seats, "seats_used": g.seats_used, "remaining": max(0, g.seats - g.seats_used)}


def _valid_uuid(v: Optional[str]) -> Optional[str]:
    try:
        return str(uuidlib.UUID(str(v))) if v else None
    except ValueError:
        return None


def scan(db: Session, user: User, token: str, count: Optional[int], request=None,
         client_uuid: Optional[str] = None) -> dict:
    from api.tickets import _can_operate_event  # même contrôle que les billets (agent assigné / orga propriétaire)
    g = db.query(EventGuest).filter(EventGuest.token == token).with_for_update().first()
    if not g:
        raise HTTPException(404, "Pass introuvable (faux pass ou QR mal lu).")
    ev = db.query(Event).filter(Event.id == g.event_id).first()
    if not _can_operate_event(user, ev):
        log_security_event(db, "scan_forbidden", request, user.id, "Pass invité %d (event %d)" % (g.id, g.event_id))
        raise HTTPException(403, "Ce pass appartient à un autre événement.")
    require_active(db, ev.id)
    cid = _valid_uuid(client_uuid)
    if cid and db.query(GuestCheckin.id).filter(GuestCheckin.client_uuid == cid).first():
        # Tentative déjà enregistrée (réponse perdue puis renvoi) : succès idempotent
        return dict(guest_public(g), kind="invitation", event=ev.title, valid=True, duplicate=True, admitted=0,
                    message="Entrée déjà enregistrée pour %s." % g.full_name)
    info = dict(guest_public(g), kind="invitation", event=ev.title)
    if g.revoked_at:
        raise HTTPException(400, "Pass ANNULÉ par l'organisateur (%s)." % g.full_name)
    remaining = g.seats - g.seats_used
    if remaining <= 0:
        raise HTTPException(400, "Pass COMPLET : %d/%d déjà entré%s (dernière entrée à %s)."
                            % (g.seats_used, g.seats, "s" if g.seats > 1 else "", fmt_time(g.last_checkin_at)))
    if count is None:
        if remaining > 1:
            return dict(info, valid=False, needs_count=True,
                        message="%s : combien de personnes entrent ?" % g.full_name)
        count = 1
    if not 1 <= count <= remaining:
        raise HTTPException(400, "Nombre invalide : %d place(s) restante(s) sur ce pass." % remaining)
    now = utcnow()
    g.seats_used += count
    g.first_checkin_at = g.first_checkin_at or now
    g.last_checkin_at = now
    db.add(GuestCheckin(guest_id=g.id, event_id=ev.id, count=count, source="online", scanned_at=now,
                        by_user_id=user.id, client_uuid=cid))
    db.commit()
    return dict(guest_public(g), kind="invitation", event=ev.title, valid=True, admitted=count,
                message="Bienvenue %s ! %d personne%s entre%s." % (g.full_name, count, "s" if count > 1 else "",
                                                                   "nt" if count > 1 else ""))


def manual_checkin(db: Session, user: User, g: EventGuest, count: int, reason: str, request=None) -> EventGuest:
    db.refresh(g, with_for_update=True)
    if g.revoked_at:
        raise HTTPException(400, "Pass annulé.")
    remaining = g.seats - g.seats_used
    if not 1 <= count <= remaining:
        raise HTTPException(400, "Nombre invalide : %d place(s) restante(s)." % remaining)
    now = utcnow()
    g.seats_used += count
    g.first_checkin_at = g.first_checkin_at or now
    g.last_checkin_at = now
    db.add(GuestCheckin(guest_id=g.id, event_id=g.event_id, count=count, source="manual", scanned_at=now,
                        by_user_id=user.id, reason=reason[:200]))
    db.commit()
    log_security_event(db, "invitation_manual_checkin", request, user.id,
                       "pass %d (event %d) +%d : %s" % (g.id, g.event_id, count, reason[:150]))
    return g


# ── hors ligne ──────────────────────────────────────────────────────
def offline_list(db: Session, ev: Event) -> dict:
    guests = db.query(EventGuest).filter(EventGuest.event_id == ev.id, EventGuest.token_hash.isnot(None)).all()
    return {
        "event": {"id": ev.id, "title": ev.title, "date": ev.event_date.isoformat() if ev.event_date else None},
        "generated_at": utcnow().isoformat() + "Z",
        "guests": [{"h": g.token_hash, "n": g.full_name, "t": g.table_label, "c": g.category,
                    "s": g.seats, "u": g.seats_used, "r": bool(g.revoked_at)} for g in guests],
    }


def sync_checkins(db: Session, user: User, device_id: Optional[str], items: list, request=None) -> list:
    """Applique les entrées faites hors ligne. Idempotent (client_uuid)."""
    from api.tickets import _can_operate_event
    results = []
    for it in items:
        raw = str(it.get("uuid") or "")
        cid = _valid_uuid(raw)
        if not cid:
            results.append({"uuid": raw, "status": "invalid"}); continue
        if db.query(GuestCheckin.id).filter(GuestCheckin.client_uuid == cid).first():
            results.append({"uuid": raw, "status": "duplicate"}); continue
        h = str(it.get("h") or "")
        g = db.query(EventGuest).filter(EventGuest.token_hash == h).with_for_update().first() if len(h) == 64 else None
        if not g:
            results.append({"uuid": raw, "status": "unknown"}); continue
        ev = db.query(Event).filter(Event.id == g.event_id).first()
        if not _can_operate_event(user, ev):
            results.append({"uuid": raw, "status": "forbidden"}); continue
        try:
            count = max(1, min(int(it.get("count") or 1), MAX_SEATS_PER_PASS))
        except (TypeError, ValueError):
            count = 1
        try:
            at = naive_utc(datetime.fromisoformat(str(it.get("at")).replace("Z", "+00:00")))
        except ValueError:
            at = utcnow()
        remaining = max(0, g.seats - g.seats_used)
        applied = min(count, remaining)
        conflict = applied < count or bool(g.revoked_at)
        g.seats_used += applied
        if applied:
            g.first_checkin_at = min(filter(None, [g.first_checkin_at, at]))
            g.last_checkin_at = max(filter(None, [g.last_checkin_at, at]))
        db.add(GuestCheckin(guest_id=g.id, event_id=g.event_id, count=count, source="offline", client_uuid=cid,
                            scanned_at=at, by_user_id=user.id, device_id=(device_id or "")[:64] or None,
                            conflict=conflict,
                            reason=("pass annulé" if g.revoked_at else "places déjà utilisées") if conflict else None))
        db.commit()
        results.append({"uuid": raw, "status": "conflict" if conflict else "ok", "applied": applied})
    n_conf = sum(1 for r in results if r["status"] == "conflict")
    if items:
        log_security_event(db, "invitation_offline_sync", request, user.id,
                           "%d entrées, %d conflits (appareil %s)" % (len(items), n_conf, (device_id or "?")[:20]))
    return results


# ── packs ────────────────────────────────────────────────────────────
def pack_price(db: Session, size: int, service: str) -> int:
    from core.pricing import get_pricing
    return int(get_pricing(db, fresh=True)["invitations"]["%s_%d" % (service, size)])


def activate_pack(db: Session, ev: Event, user: User, size: int, service: str, channel: str, ref: str,
                  amount_received: int, confirmed: bool, note: Optional[str], request=None) -> InvitationPack:
    if size not in INVITATION_PACK_SIZES:
        raise HTTPException(400, "Taille de pack invalide.")
    if service not in INVITATION_SERVICE_LEVELS:
        raise HTTPException(400, "Niveau de service invalide.")
    ch = next((c for c in channels(active_only=True) if c["code"] == channel), None)
    if not ch:
        raise HTTPException(400, "Moyen de paiement non disponible.")
    if not confirmed:
        raise HTTPException(400, "Confirmez que le paiement a été reçu sur un moyen au nom de GoEvent SARLU.")
    ref = re.sub(r"\s+", " ", (ref or "").strip())
    if len(ref) < 3:
        raise HTTPException(400, "Référence du paiement obligatoire (%s)." % ch["ref_label"])
    price = pack_price(db, size, service)
    if amount_received < price:
        raise HTTPException(400, "Montant reçu (%d FCFA) inférieur au prix du pack (%d FCFA)." % (amount_received, price))
    if db.query(InvitationPack.id).filter(InvitationPack.payment_channel == channel,
                                          func.upper(InvitationPack.payment_ref) == ref.upper()).first():
        raise HTTPException(409, "Cette référence de paiement a déjà servi à activer un pack.")
    listed = seats_listed(db, ev.id)
    if quota(db, ev.id) + size < listed:
        raise HTTPException(400, "Pack trop petit : %d personnes déjà dans la liste." % listed)
    pack = InvitationPack(event_id=ev.id, size=size, service=service, amount_fcfa=price, payment_channel=channel,
                          payment_ref=ref[:80], amount_received=amount_received, confirmed_goevent_account=True,
                          activated_by=user.id, note=(note or "")[:300] or None)
    db.add(pack)
    db.commit()
    db.refresh(pack)
    log_security_event(db, "invitation_pack_activated", request, user.id,
                       "event %d : pack %d %s, %d FCFA, %s réf. %s (reçu %d)"
                       % (ev.id, size, service, price, channel, ref[:40], amount_received))
    return pack


# ── purge J+30 ───────────────────────────────────────────────────────
def purge_after_for(event_date: Optional[datetime]) -> datetime:
    return (event_date or utcnow()) + timedelta(days=INVITATION_RETENTION_DAYS)


def purge_due(db: Session, force_event_id: Optional[int] = None) -> int:
    """Anonymise les listes d'invités dont la date de conservation est passée."""
    q = db.query(PrivateEventSettings).filter(PrivateEventSettings.purged_at.is_(None))
    q = q.filter(PrivateEventSettings.event_id == force_event_id) if force_event_id else \
        q.filter(PrivateEventSettings.purge_after <= utcnow())
    n = 0
    for st in q.all():
        n += db.query(EventGuest).filter(EventGuest.event_id == st.event_id, EventGuest.purged == False).update(  # noqa: E712
            {EventGuest.full_name: "Invité supprimé", EventGuest.phone: None, EventGuest.token: None,
             EventGuest.token_hash: None, EventGuest.purged: True}, synchronize_session=False)
        db.query(GuestCheckin).filter(GuestCheckin.event_id == st.event_id).update(
            {GuestCheckin.device_id: None, GuestCheckin.reason: None}, synchronize_session=False)
        st.client_phone = None
        st.track_token = None
        st.purged_at = utcnow()
    db.commit()
    return n
